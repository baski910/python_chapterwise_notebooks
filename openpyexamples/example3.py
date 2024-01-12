from openpyxl import Workbook
from openpyxl.chart import BarChart,Series,Reference

wb=Workbook()
ws=wb.create_sheet()

rows=[('Number','Batch 1','Batch 2'),
      (2, 10, 30),
    (3, 40, 60),
    (4, 50, 70),
    (5, 20, 10),
    (6, 10, 40),
    (7, 50, 30),
]
for row in rows:
    ws.append(row)

chart1 = BarChart()
chart1.type = "col"
chart1.style = 10
chart1.title = "Bar Chart"
chart1.y_axis.title = 'Test number'
chart1.x_axis.title = 'Sample length (mm)'

data=Reference(ws,min_col=2,min_row=1,max_row=7,max_col=3)
cats=Reference(ws,min_col=1,min_row=2,max_row=7)
chart1.add_data(data,titles_from_data=True)
chart1.set_categories(cats)
ws.add_chart(chart1,"A10")

wb.save('chart_example.xlsx')
